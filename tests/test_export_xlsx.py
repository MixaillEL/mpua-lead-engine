import io
from datetime import datetime

from openpyxl import load_workbook

from exporters.models import ExportRow
from exporters.xlsx_exporter import build_workbook


def _row(**overrides) -> ExportRow:
    defaults = dict(
        company_id="c1",
        company_name="Test Co",
        website="https://example.com",
        email_1="info@example.com",
        first_seen=datetime(2026, 1, 1, 10, 0),
        last_seen=datetime(2026, 2, 1, 10, 0),
    )
    defaults.update(overrides)
    return ExportRow(**defaults)


def _summary(**overrides) -> dict:
    defaults = dict(
        job_id="job-1",
        job_run_id="run-1",
        preset="dentist",
        regions="Kyiv",
        status="completed",
        unique_companies=1,
        with_phone=1,
        with_email=1,
        with_website=1,
        with_social=0,
        osm_attribution_required=False,
    )
    defaults.update(overrides)
    return defaults


def _load(rows, summary):
    wb = build_workbook(rows, summary)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


def test_workbook_opens_and_has_expected_sheets():
    wb = _load([_row()], _summary())
    assert "Companies" in wb.sheetnames
    assert "Summary" in wb.sheetnames


def test_companies_row_count_matches_input():
    rows = [_row(company_id=f"c{i}", company_name=f"Company {i}") for i in range(3)]
    wb = _load(rows, _summary())
    ws = wb["Companies"]
    assert ws.max_row == 4  # header + 3


def test_freeze_panes_set():
    wb = _load([_row()], _summary())
    assert wb["Companies"].freeze_panes == "A2"


def test_autofilter_set():
    wb = _load([_row()], _summary())
    assert wb["Companies"].auto_filter.ref is not None


def test_hyperlinks_present_for_website_and_email():
    wb = _load([_row()], _summary())
    ws = wb["Companies"]
    header = [c.value for c in ws[1]]
    website_col = header.index("Website") + 1
    email_col = header.index("Email 1") + 1

    website_cell = ws.cell(row=2, column=website_col)
    email_cell = ws.cell(row=2, column=email_col)

    assert website_cell.hyperlink is not None
    assert email_cell.hyperlink is not None
    assert email_cell.hyperlink.target.startswith("mailto:")


def test_ukrainian_text_preserved():
    wb = _load([_row(company_name="Клімат Сервіс", city="Дніпро")], _summary())
    ws = wb["Companies"]
    header = [c.value for c in ws[1]]
    name_col = header.index("Company") + 1
    assert ws.cell(row=2, column=name_col).value == "Клімат Сервіс"


def test_summary_metrics_present():
    wb = _load([_row()], _summary(unique_companies=42, with_phone=10))
    ws = wb["Summary"]
    values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert values.get("Companies") == 42
    assert values.get("With Phone") == 10


def test_empty_job_run_produces_header_only_sheet():
    wb = _load([], _summary(unique_companies=0))
    ws = wb["Companies"]
    assert ws.max_row == 1  # header only


def test_osm_attribution_note_present_when_required():
    wb = _load([_row()], _summary(osm_attribution_required=True))
    ws = wb["Summary"]
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "OpenStreetMap" in all_text
    assert "ODbL" in all_text


def test_osm_attribution_note_absent_when_not_required():
    wb = _load([_row()], _summary(osm_attribution_required=False))
    ws = wb["Summary"]
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "OpenStreetMap" not in all_text


def test_date_columns_formatted():
    wb = _load([_row()], _summary())
    ws = wb["Companies"]
    header = [c.value for c in ws[1]]
    first_seen_col = header.index("First Seen") + 1
    cell = ws.cell(row=2, column=first_seen_col)
    assert cell.number_format != "General"
