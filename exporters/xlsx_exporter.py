"""XLSX export via openpyxl: a `Companies` sheet (one row per Company)
plus a `Summary` sheet built from JobRun.metrics — never recomputed
separately, so it can't drift from what the orchestrator persisted.
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from exporters.csv_exporter import COLUMNS, sanitize_cell
from exporters.models import ExportRow

HEADER_FONT = Font(bold=True)
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")

COLUMN_WIDTHS = {
    "Company": 32,
    "Category": 16,
    "Country": 10,
    "Region": 14,
    "City": 14,
    "Address": 36,
    "Phone 1": 16,
    "Phone 2": 16,
    "Phone 3": 16,
    "Email 1": 26,
    "Email 2": 26,
    "Email 3": 26,
    "Website": 30,
    "Facebook": 26,
    "Instagram": 26,
    "Telegram": 22,
    "LinkedIn": 26,
    "YouTube": 26,
    "TikTok": 22,
    "Rating": 8,
    "Reviews": 9,
    "Source Types": 22,
    "Source URLs": 40,
    "First Seen": 18,
    "Last Seen": 18,
}

WRAP_COLUMNS = {"Address", "Source URLs"}
DATE_COLUMNS = {"First Seen", "Last Seen"}
HYPERLINK_COLUMNS = {
    "Website": "url",
    "Email 1": "mailto",
    "Email 2": "mailto",
    "Email 3": "mailto",
}


def _write_companies_sheet(ws: Worksheet, rows: list[ExportRow]) -> None:
    headers = [header for header, _ in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for row in rows:
        data = row.model_dump()
        values = []
        for header, field in COLUMNS:
            value = data.get(field)
            if field not in ("rating", "reviews"):
                value = sanitize_cell(value, field)
            values.append(value)
        ws.append(values)

    last_row = ws.max_row
    last_col = len(headers)

    for col_idx, (header, field) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(header, 18)

        if header in DATE_COLUMNS:
            for row_idx in range(2, last_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"

        if header in WRAP_COLUMNS:
            for row_idx in range(2, last_row + 1):
                ws.cell(row=row_idx, column=col_idx).alignment = WRAP_ALIGNMENT

        if header in HYPERLINK_COLUMNS:
            kind = HYPERLINK_COLUMNS[header]
            for row_idx in range(2, last_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value and isinstance(cell.value, str) and not cell.value.startswith("'"):
                    cell.hyperlink = cell.value if kind == "url" else f"mailto:{cell.value}"
                    cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    if last_row >= 1 and last_col >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{max(last_row, 1)}"


SUMMARY_FIELDS = [
    ("Job ID", "job_id"),
    ("JobRun ID", "job_run_id"),
    ("Preset", "preset"),
    ("Regions", "regions"),
    ("Status", "status"),
    ("Started", "started_at"),
    ("Finished", "finished_at"),
    ("Runtime (s)", "runtime_seconds"),
    ("Companies", "unique_companies"),
    ("With Phone", "with_phone"),
    ("With Email", "with_email"),
    ("With Website", "with_website"),
    ("With Social", "with_social"),
    ("Phone Coverage", "phone_coverage_pct"),
    ("Email Coverage", "email_coverage_pct"),
    ("Website Coverage", "website_coverage_pct"),
    ("Social Coverage", "social_coverage_pct"),
    ("New Companies", "new_companies"),
    ("Matched Companies", "matched"),
    ("Review Count", "review"),
    ("OSM Candidates", "osm_candidates"),
    ("Tavily Candidates", "tavily_candidates"),
    ("Tavily Requests", "api_requests"),
    ("External Cost USD", "api_cost_usd"),
    ("Exported At", "exported_at"),
]


def _write_summary_sheet(ws: Worksheet, summary: dict) -> None:
    ws.append(["Field", "Value"])
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for label, key in SUMMARY_FIELDS:
        ws.append([label, summary.get(key)])

    if summary.get("osm_attribution_required"):
        ws.append([])
        ws.append(
            [
                "Some records may contain information from OpenStreetMap, "
                "available under the Open Database License (ODbL)."
            ]
        )

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40


def build_workbook(rows: list[ExportRow], summary: dict) -> Workbook:
    wb = Workbook()
    companies_ws = wb.active
    companies_ws.title = "Companies"
    _write_companies_sheet(companies_ws, rows)

    summary_ws = wb.create_sheet("Summary")
    _write_summary_sheet(summary_ws, summary)

    return wb
